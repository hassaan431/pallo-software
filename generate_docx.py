import os
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from database import get_public_products

def generate_brochure():
    products = get_public_products()
    brands = {}
    brand_logos = {}
    
    # Group by brand
    from database import get_brand_logo_filename
    for p in products:
        b = p.get('brand') or 'Other'
        if b not in brands:
            brands[b] = []
            logo_file = get_brand_logo_filename(b)
            brand_logos[b] = logo_file
        brands[b].append(p)
        
    doc = Document()
    
    # Main logo (Pushed down)
    for _ in range(5):
        doc.add_paragraph()
        
    main_logo_path = os.path.join('static', 'logos', 'AK Distributors Logo.jpeg')
    if os.path.exists(main_logo_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(main_logo_path, width=Inches(3.0))
    
    # Cover Page Title
    title = doc.add_heading('AK Distributors Catalog', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Push brands to bottom
    for _ in range(8):
        doc.add_paragraph()
        
    p = doc.add_paragraph("Our Partner Brands")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p2.add_run()
    
    for b, logo_file in brand_logos.items():
        if logo_file:
            logo_path = os.path.join('static', 'logos', 'brands', logo_file)
            run.add_picture(logo_path, width=Inches(0.6))
            run.add_text("  ")
            
    # Add page break
    doc.add_page_break()
    
    for brand, brand_products in brands.items():
        # Brand Heading
        p = doc.add_heading('', level=1)
        r = p.add_run()
        if brand_logos.get(brand):
            logo_path = os.path.join('static', 'logos', 'brands', brand_logos[brand])
            try:
                r.add_picture(logo_path, width=Inches(0.3))
                r.add_text("   ")
            except:
                pass
        r.add_text(brand)
        
        # Product Table (2 columns)
        table = doc.add_table(rows=0, cols=2)
        table.style = 'Table Grid'
        
        for i in range(0, len(brand_products), 2):
            row_cells = table.add_row().cells
            
            # First product
            p1 = brand_products[i]
            para1 = row_cells[0].paragraphs[0]
            para1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if p1.get('image_url'):
                filename = p1['image_url'].split('/')[-1]
                img_path = os.path.join('static', 'uploads', 'products', filename)
                if os.path.exists(img_path):
                    p_run = para1.add_run()
                    try:
                        p_run.add_picture(img_path, width=Inches(1.5))
                    except:
                        pass
            name_para1 = row_cells[0].add_paragraph(p1['name'])
            name_para1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Second product (if exists)
            if i + 1 < len(brand_products):
                p2 = brand_products[i+1]
                para2 = row_cells[1].paragraphs[0]
                para2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if p2.get('image_url'):
                    filename = p2['image_url'].split('/')[-1]
                    img_path = os.path.join('static', 'uploads', 'products', filename)
                    if os.path.exists(img_path):
                        p_run = para2.add_run()
                        try:
                            p_run.add_picture(img_path, width=Inches(1.5))
                        except:
                            pass
                name_para2 = row_cells[1].add_paragraph(p2['name'])
                name_para2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_page_break()
        
    doc.save('AK_Distributors_Brochure_Draft.docx')
    print("Draft generated: AK_Distributors_Brochure_Draft.docx")

if __name__ == '__main__':
    generate_brochure()
